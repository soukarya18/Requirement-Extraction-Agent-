from __future__ import annotations

import base64
import importlib
import mimetypes
import os
import threading
import uuid
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from dash import Dash, Input, Output, State, ALL, ctx, dcc, html, no_update
from dash.exceptions import PreventUpdate
from flask import send_file

import webbrowser
from threading import Timer
from dash import Dash , html 

# ---------------------------------------------------------------------------
# Paths / configuration
# ---------------------------------------------------------------------------

BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in os.sys.path:
    os.sys.path.insert(0, str(BASE_DIR))

UPLOADS_DIR = BASE_DIR / "data" / "uploads"
OUTPUT_DIR = BASE_DIR / "output"
SUPPORTED_EXTENSIONS = {"pdf", "docx", "doc", "txt"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# ---------------------------------------------------------------------------
# Thread-safe background job registry
# ---------------------------------------------------------------------------

_JOB_LOCK = threading.Lock()
_JOB_RESULTS: Dict[Tuple[str, str], Dict[str, Any]] = {}


def _publish_job_result(session_uid: str, job_name: str, *, status: str,
                        result: Any = None, error: Optional[str] = None) -> None:
    with _JOB_LOCK:
        _JOB_RESULTS[(session_uid, job_name)] = {
            "status": status,
            "result": result,
            "error": error,
        }


def _start_background_job(session_uid: str, job_name: str, target: Any,
                          kwargs: Dict[str, Any]) -> bool:
    key = (session_uid, job_name)
    with _JOB_LOCK:
        if key in _JOB_RESULTS and _JOB_RESULTS[key].get("status") == "running":
            return False
        _JOB_RESULTS[key] = {"status": "running", "result": None, "error": None}

    def _runner() -> None:
        try:
            result = target(**kwargs)
            _publish_job_result(session_uid, job_name, status="done", result=result)
        except Exception as exc:  # pragma: no cover - backend-defined failures
            _publish_job_result(session_uid, job_name, status="error", error=str(exc))

    worker = threading.Thread(
        target=_runner,
        name=f"sdlc_{job_name.replace(':', '_')}_worker",
        daemon=True,
    )
    worker.start()
    return True


def _pending_jobs(session_uid: str) -> List[str]:
    with _JOB_LOCK:
        return [
            job_name
            for (uid, job_name), payload in _JOB_RESULTS.items()
            if uid == session_uid and payload.get("status") == "running"
        ]


def _drain_session_results(session_uid: str) -> List[Tuple[str, Dict[str, Any]]]:
    drained: List[Tuple[str, Dict[str, Any]]] = []
    with _JOB_LOCK:
        for (uid, job_name), payload in list(_JOB_RESULTS.items()):
            if uid != session_uid or payload.get("status") == "running":
                continue
            drained.append((job_name, payload))
            del _JOB_RESULTS[(uid, job_name)]
    return drained

# ---------------------------------------------------------------------------
# Backend adapters
# ---------------------------------------------------------------------------


def _backend_callable(module_name: str, function_name: str) -> Any:
    module = importlib.import_module(module_name)
    return getattr(module, function_name)


def _save_upload(contents: str, filename: str) -> Path:
    suffix = Path(filename).suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: .{suffix or 'unknown'}")
    if not contents or "," not in contents:
        raise ValueError("The uploaded file payload is invalid.")

    _header, encoded = contents.split(",", 1)
    raw = base64.b64decode(encoded)
    UPLOADS_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = Path(filename).name
    destination = UPLOADS_DIR / safe_name
    destination.write_bytes(raw)
    return destination


def _result_path(result: Any) -> Optional[Path]:
    value = getattr(result, "generated_document", None)
    if not value:
        return None
    path = Path(str(value))
    return path if path.is_file() else None


def _result_markdown(result: Any) -> Optional[Path]:
    value = getattr(result, "generated_markdown", None)
    if not value:
        return None
    path = Path(str(value))
    return path if path.is_file() else None


def _docx_to_markdown(docx_path: Path, output_path: Path) -> Path:
    """Mirror the Streamlit DOCX->Markdown conversion used for generated artifacts."""
    from docx import Document

    document = Document(str(docx_path))
    lines: List[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        if "title" in style_name or "heading 1" in style_name:
            lines.append(f"# {text}")
        elif "heading 2" in style_name:
            lines.append(f"## {text}")
        elif "heading 3" in style_name:
            lines.append(f"### {text}")
        elif text.startswith("• "):
            lines.append(f"- {text[2:]}")
        elif text.startswith("- "):
            lines.append(text)
        else:
            lines.append(text)

    for table in document.tables:
        if not table.rows:
            continue
        lines.append("")
        rows = [[cell.text.replace("\n", " ").strip() for cell in row.cells]
                for row in table.rows]
        if rows:
            lines.append("| " + " | ".join(rows[0]) + " |")
            lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
            for row in rows[1:]:
                lines.append("| " + " | ".join(row) + " |")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n\n".join(lines).strip() + "\n", encoding="utf-8")
    return output_path


def _safe_existing_path(raw_path: str) -> Optional[Path]:
    """Allow downloads only from this app's upload/output workspaces."""
    if not raw_path:
        return None
    try:
        path = Path(raw_path).resolve()
        allowed_roots = [UPLOADS_DIR.resolve(), OUTPUT_DIR.resolve()]
        if any(root == path or root in path.parents for root in allowed_roots):
            return path if path.is_file() else None
    except (OSError, ValueError):
        pass
    return None


def _find_brs_source(brs_name: str) -> Optional[Path]:
    if not UPLOADS_DIR.is_dir():
        return None
    for candidate in UPLOADS_DIR.iterdir():
        if candidate.is_file() and candidate.stem == brs_name and candidate.suffix.lower() == ".docx":
            return candidate
    return None


def _find_estimation_artifact(folder: Path) -> Optional[Path]:
    if not folder.is_dir():
        return None
    candidates = sorted(
        [
            path for path in folder.iterdir()
            if path.is_file()
            and path.suffix.lower() == ".xlsx"
            and path.name.lower() == "generated_estimation.xlsx"
        ],
        key=lambda path: path.name.lower(),
    )
    return candidates[0] if candidates else None


def _discover_dashboard_rows() -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if not OUTPUT_DIR.is_dir():
        return rows

    folders = sorted(
        [path for path in OUTPUT_DIR.iterdir() if path.is_dir()],
        key=lambda path: path.name.lower(),
    )
    for folder in folders:
        brs_name = folder.name
        brd_path = _find_brs_source(brs_name)
        drs_path = folder / "generated_drs.docx"
        jira_path = folder / "generated_jira_user_stories.docx"
        estimation_path = _find_estimation_artifact(folder)
        files = (brd_path, drs_path, jira_path, estimation_path)
        if not any(path and path.is_file() for path in files):
            continue
        rows.append({
            "Title": brs_name,
            "BRD": brd_path,
            "DRS": drs_path if drs_path.is_file() else None,
            "Jira": jira_path if jira_path.is_file() else None,
            "Estimation": estimation_path,
        })
    return rows


def _extract_kb_path(result: Any) -> Optional[str]:
    """Compatibility hook for ingestion results that expose a KB path.

    The uploaded Streamlit UI defines chat_kb_path but never assigns it directly.
    We therefore only populate it when the ingestion result exposes a conventional
    path-like attribute; otherwise the chatbot reports the backend's own error.
    """
    for attr in ("knowledge_base_path", "kb_path", "knowledge_base"):
        value = getattr(result, attr, None)
        if value:
            return str(value)
    return result if isinstance(result, (str, os.PathLike)) else None


# ---------------------------------------------------------------------------
# Application state
# ---------------------------------------------------------------------------

DEFAULT_STATE: Dict[str, Any] = {
    "active_module": "requirements",
    "uploaded_brd": None,
    "ingestion_status": "idle",
    "drs_status": "idle",
    "jira_status": "idle",
    "estimation_status": "idle",
    "ingestion_error": None,
    "drs_error": None,
    "jira_error": None,
    "estimation_error": None,
    "chat_kb_path": None,
    "drs_word": None,
    "drs_md": None,
    "jira_word": None,
    "jira_md": None,
    "estimation_xlsx": None,
    "chat_history": [],
    "chat_job": None,
}


def fresh_state() -> Dict[str, Any]:
    return dict(DEFAULT_STATE)

# ---------------------------------------------------------------------------
# UI helpers
# ---------------------------------------------------------------------------


def button(label: str, component_id: str, *, kind: str = "secondary", disabled: bool = False):
    classes = f"btn btn-{kind}"
    return html.Button(label, id=component_id, className=classes, disabled=disabled)


def status_banner(message: str, tone: str = "info"):
    return html.Div(message, className=f"status-banner status-{tone}")


def artifact_card(title: str, status: str, path: Optional[str], download_id: str, label: str,
                  output_type: str = "generic"):
    ready = bool(path and Path(path).is_file())
    children = [
        html.Div(title, className="artifact-name"),
        html.Div(
            [
                html.Button(
                    "Ready — Download Excel",
                    id={"type": "estimation-download", "path": str(path)},
                    className="artifact-ready-download",
                )
                if ready and output_type == "estimation"
                else (html.Span("Ready", className="ready") if ready else status)
            ],
            className="artifact-status",
        ),
    ]
    if ready and output_type != "estimation":
        children.append(button(label, download_id, kind="secondary"))
    elif not ready:
        children.append(html.Div("Not generated yet", className="artifact-empty"))
    return html.Div(children, className=f"artifact-card {output_type}")


def chat_bubble(message: Dict[str, Any], index: int):
    role = message.get("role", "assistant")
    content = message.get("content", "")
    pending = message.get("pending") is True
    return html.Div(
        [
            html.Div("You" if role == "user" else "Assistant", className="chat-role"),
            dcc.Markdown(content, className="chat-content"),
            html.Div("Searching the knowledge base...", className="chat-pending") if pending else None,
        ],
        className=f"chat-bubble chat-{role}",
        key=f"chat-{index}",
    )


def nav_panel(state: Dict[str, Any]):
    active = state.get("active_module", "requirements")
    return html.Div(
        [
            html.Div("WORKSPACE", className="nav-label"),
            html.Button(
                "Dashboard", id="nav-dashboard",
                className=f"nav-button {'active' if active == 'dashboard' else ''}",
            ),
            html.Button(
                "Requirements Generation", id="nav-requirements",
                className=f"nav-button {'active' if active == 'requirements' else ''}",
            ),
            html.Button(
                "Estimation Generation", id="nav-estimation",
                className=f"nav-button {'active' if active == 'estimation' else ''}",
            ),
            html.Div(
                "Create module-level person-day estimates from the processed BRD.",
                className="small-note nav-note",
            ),
        ],
        className="panel-card nav-panel",
    )


def requirements_view(state: Dict[str, Any]):
    ingestion = state.get("ingestion_status", "idle")
    ready = ingestion == "done"
    running_any = any(
        state.get(f"{name}_status") == "running"
        for name in ("ingestion", "drs", "jira", "estimation")
    )

    if ingestion == "running":
        ingestion_message = status_banner("Building the active knowledge base...", "info")
    elif ingestion == "done":
        ingestion_message = status_banner("BRD processed successfully. You can now generate the required artifacts.", "success")
    elif ingestion == "error":
        ingestion_message = status_banner(
            f"BRD processing failed: {state.get('ingestion_error') or 'Unknown error'}", "error"
        )
    else:
        ingestion_message = status_banner("Upload a BRD to begin.", "muted")

    drs_status = state.get("drs_status", "idle")
    jira_status = state.get("jira_status", "idle")

    return html.Div(
        [
            html.Div("REQUIREMENT GENERATION", className="module-caption"),
            html.H1("SDLC Lifecycle Assistant", className="portal-title"),
            html.P(
                "Upload a BRD, generate DRS/Jira artifacts, and ask evidence-grounded questions.",
                className="portal-subtitle",
            ),
            html.Div(
                [
                    html.Div(
                        [
                            html.Div("BRD Upload", className="section-title"),
                            dcc.Upload(
                                id="brd-upload",
                                children=html.Div([
                                    html.Div("Drop your BRD here", className="upload-title"),
                                    html.Div("or click to browse · PDF / DOCX / DOC / TXT", className="upload-subtitle"),
                                ]),
                                multiple=False,
                                className="upload-zone",
                            ),
                            html.Div(id="upload-file-name", className="upload-file-name"),
                            button("Process BRD", "process-brd", kind="primary", disabled=running_any),
                        ],
                        className="upload-block",
                    ),
                    html.Div(id="ingestion-banner", children=ingestion_message, className="full-width"),
                    html.Div(
                        [
                            html.Div(
                                [
                                    button(
                                        "Generating DRS..." if drs_status == "running" else "Generate DRS",
                                        "generate-drs", kind="primary", disabled=(not ready or drs_status == "running"),
                                    )
                                ], className="action-col"
                            ),
                            html.Div(
                                [
                                    button(
                                        "Generating Jira User Stories..." if jira_status == "running" else "Generate Jira User Stories",
                                        "generate-jira", kind="primary", disabled=(not ready or jira_status == "running"),
                                    )
                                ], className="action-col"
                            ),
                        ],
                        className="action-row",
                    ) if ready else html.Div("Process the BRD to enable DRS and Jira generation.", className="disabled-hint"),
            ], className="panel-card primary-work-card"),
            html.Div(
                [
                    html.Div(
                        [
                            html.Span("Source", className="meta-label"),
                            html.Span(Path(state["uploaded_brd"]).name if state.get("uploaded_brd") else "No BRD selected", className="meta-value"),
                        ], className="meta-block",
                    ),
                    html.Div(
                        [
                            html.Span("Knowledge base", className="meta-label"),
                            html.Span("Ready" if ready else "Not ready", className=f"meta-value {'ready' if ready else ''}"),
                        ], className="meta-block",
                    ),
                ], className="meta-row",
            ),
            html.Div(
                [
                    html.Div("QA Chatbot", className="section-title"),
                    html.Div(
                        "Ask a question about the uploaded BRD. Answers are grounded in the active knowledge base.",
                        className="small-note",
                    ) if not state.get("chat_history") else None,
                    html.Div(
                        [chat_bubble(m, i) for i, m in enumerate(state.get("chat_history", []))]
                        or [status_banner("Process a BRD before asking questions.", "muted")],
                        id="chat-history",
                        className="chat-history",
                    ),
                    html.Div(
                        [
                            dcc.Input(
                                id="chat-input",
                                placeholder="Ask about requirements, rules, processes, NFRs, or traceability...",
                                type="text",
                                className="chat-input",
                                disabled=(not ready or running_any or bool(state.get("chat_job"))),
                            ),
                            button("Send", "chat-send", kind="primary", disabled=(not ready or running_any or bool(state.get("chat_job")))),
                        ], className="chat-composer",
                    ),
                    status_banner(
                        "Chatbot is disabled while a BRD, DRS, Jira, or estimation process is running.", "info"
                    ) if running_any and ready else None,
                ],
                className="panel-card chat-card",
            ),
            dcc.Interval(id="chat-scroll-trigger", interval=250, n_intervals=0),
        ],
        className="workspace-stack",
    )


def estimation_view(state: Dict[str, Any]):
    ready = state.get("ingestion_status") == "done"
    status = state.get("estimation_status", "idle")
    if status == "running":
        banner = status_banner("Extracting BRD evidence and calculating the weighted complexity score...", "info")
    elif status == "done":
        banner = status_banner("Estimation generation completed.", "success")
    elif status == "error":
        banner = status_banner(f"Estimation generation failed: {state.get('estimation_error') or 'Unknown error'}", "error")
    else:
        banner = status_banner("Process a BRD in Requirement Generation before generating an estimation.", "muted")

    return html.Div(
        [
            html.Div("ESTIMATION GENERATION", className="module-caption"),
            html.H1("SDLC Lifecycle Assistant", className="portal-title"),
            html.P(
                "Generate an evidence-grounded complexity estimation from the processed BRD.",
                className="portal-subtitle",
            ),
            html.Div(
                [
                    html.Div(
                        "Generate Estimation",
                        className="section-title",
                    ),
                    button(
                        "Generating Estimation..." if status == "running" else "Generate Estimation",
                        "generate-estimation", kind="primary",
                        disabled=(not ready or status == "running"),
                    ),
                    dcc.Checklist(
                        id="qa-support",
                        options=[{"label": "QA support is provided", "value": "included"}],
                        value=[],
                        className="qa-support-checklist",
                    ),
                    html.Div(className="spacer"),
                    banner,
                ],
                className="panel-card estimation-card",
            ),
        ],
        className="workspace-stack",
    )


def dashboard_view():
    rows = _discover_dashboard_rows()
    if not rows:
        body = status_banner("No generated BRS DOCX artefacts found in the output directory yet.", "muted")
    else:
        headers = ["Sl no", "Title", "BRD Link", "DRS Link", "Jira link", "Estimation"]
        header = html.Div(
            [html.Div(h, className="grid-header-cell") for h in headers],
            className="dashboard-grid dashboard-header",
        )
        rendered_rows = []
        for idx, row in enumerate(rows, start=1):
            cells = [
                html.Div(str(idx), className="grid-cell grid-cell-index"),
                html.Div(row["Title"], className="grid-cell grid-cell-title", title=row["Title"]),
            ]
            for key in ("BRD", "DRS", "Jira", "Estimation"):
                path = row.get(key)
                if path and Path(path).is_file():
                    kind = f"dashboard-{idx}-{key.lower()}"
                    cells.append(
                        html.Div(
                            [
                                html.Button(
                                    "Download",
                                    id={"type": "dashboard-download", "kind": kind, "path": str(path)},
                                    className="btn btn-secondary btn-small",
                                )
                            ],
                            className="grid-cell",
                        )
                    )
                else:
                    cells.append(html.Div("—", className="grid-cell grid-cell-empty"))
            rendered_rows.append(html.Div(cells, className="dashboard-grid dashboard-row"))
        body = html.Div([header] + rendered_rows, className="dashboard-table")

    return html.Div(
        [
            html.Div("DASHBOARD", className="module-caption"),
            html.H1("SDLC Dashboard", className="portal-title"),
            html.P(
                "Persisted DOCX artefacts discovered from the output workspace.",
                className="portal-subtitle",
            ),
            html.Div(body, className="panel-card dashboard-card"),
        ],
        className="workspace-stack",
    )


def output_panel(state: Dict[str, Any]):
    module = state.get("active_module", "requirements")
    if module == "dashboard":
        return html.Div(style={"display": "none"})

    if module == "requirements":
        content = [
            html.Div("DRS Documents", className="output-subheading"),
            artifact_card("DRS Word", "Not generated", state.get("drs_word"), "download-drs-word", "Download Word"),
            artifact_card("DRS Markdown", "Not generated", state.get("drs_md"), "download-drs-md", "Download Markdown"),
            html.Div("Jira User Stories", className="output-subheading"),
            artifact_card("Jira Stories Word", "Not generated", state.get("jira_word"), "download-jira-word", "Download Word"),
            artifact_card("Jira Stories Markdown", "Not generated", state.get("jira_md"), "download-jira-md", "Download Markdown"),
        ]
    else:
        content = [
            html.Div("Estimation Report", className="output-subheading"),
            artifact_card("Estimation Excel", "Not generated", state.get("estimation_xlsx"), "download-estimation-xlsx", "Download Excel", output_type="estimation"),
        ]

    return html.Div(
        [html.Div("OUTPUT & GENERATED ARTIFACTS", className="section-title")] + content,
        className="panel-card output-panel",
    )


def serve_layout():
    session_uid = uuid.uuid4().hex
    state = fresh_state()
    return html.Div(
        [
            dcc.Store(id="session-store", data={"session_uid": session_uid}),
            dcc.Store(id="app-state", data=state),
            # Poll only while a background job is active. This avoids an
            # always-on stream of /_dash-update-component requests when the
            # application is idle.
            dcc.Interval(id="job-poller", interval=2000, n_intervals=0, disabled=True),
            dcc.Download(id="file-download"),
            html.Div(
                [
                    html.Div("SDLC Assistant Portal", className="portal-title page-title"),
                    html.Div("AI-powered requirement generation and lifecycle assistance", className="portal-subtitle"),
                ],
                className="page-header",
            ),
            html.Div(
                [
                    html.Div(id="left-nav"),
                    html.Div(id="main-workspace"),
                    html.Div(id="right-output"),
                ],
                className="app-grid",
            ),
        ],
        className="app-shell",
    )


app = Dash(__name__, title="SDLC Assistant Portal", suppress_callback_exceptions=True)
app.layout = serve_layout

# ---------------------------------------------------------------------------
# Navigation + shell rendering
# ---------------------------------------------------------------------------


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("nav-dashboard", "n_clicks"),
    Input("nav-requirements", "n_clicks"),
    Input("nav-estimation", "n_clicks"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def set_active_module(_dashboard, _requirements, _estimation, state):
    triggered = ctx.triggered_id
    if not triggered:
        raise PreventUpdate

    # These navigation buttons are recreated whenever app-state changes.
    # Dash can invoke this callback when the dynamic inputs are inserted, even
    # though the user did not click anything. Only a positive n_clicks value is
    # a real navigation action.
    click_counts = {
        "nav-dashboard": _dashboard,
        "nav-requirements": _requirements,
        "nav-estimation": _estimation,
    }
    if not click_counts.get(triggered):
        raise PreventUpdate

    state = dict(state or fresh_state())
    mapping = {
        "nav-dashboard": "dashboard",
        "nav-requirements": "requirements",
        "nav-estimation": "estimation",
    }
    state["active_module"] = mapping[triggered]
    return state


@app.callback(
    Output("left-nav", "children"),
    Output("main-workspace", "children"),
    Output("right-output", "children"),
    Output("job-poller", "disabled"),
    Input("app-state", "data"),
)
def render_shell(state):
    state = state or fresh_state()
    nav = nav_panel(state)
    active = state.get("active_module", "requirements")
    if active == "dashboard":
        workspace = dashboard_view()
    elif active == "estimation":
        workspace = estimation_view(state)
    else:
        workspace = requirements_view(state)

    # Keep polling enabled only while work that needs completion detection is
    # actually running. Once the job completes, the next state update disables
    # the interval again.
    polling_needed = any(
        state.get(f"{name}_status") == "running"
        for name in ("ingestion", "drs", "jira", "estimation")
    ) or bool(state.get("chat_job"))

    return nav, workspace, output_panel(state), not polling_needed


# ---------------------------------------------------------------------------
# Upload / processing
# ---------------------------------------------------------------------------


@app.callback(
    Output("upload-file-name", "children"),
    Input("brd-upload", "filename"),
)
def show_uploaded_filename(filename):
    if not filename:
        return ""
    return html.Div(f"Selected: {filename}", className="upload-file-name")


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("process-brd", "n_clicks"),
    State("brd-upload", "contents"),
    State("brd-upload", "filename"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def process_brd(_n_clicks, contents, filename, session_data, state):
    if not contents or not filename:
        raise PreventUpdate
    state = dict(state or fresh_state())
    session_uid = session_data["session_uid"]
    try:
        saved_path = _save_upload(contents, filename)
        # Reset generated artifacts for a new source document.
        reset = fresh_state()
        reset["active_module"] = state.get("active_module", "requirements")
        reset["uploaded_brd"] = str(saved_path)
        reset["ingestion_status"] = "running"
        _start_background_job(
            session_uid,
            "ingestion",
            _backend_callable("src.ingestion_pipeline", "ingest_brd"),
            {"brd_path": saved_path},
        )
        return reset
    except Exception as exc:
        state["ingestion_status"] = "error"
        state["ingestion_error"] = str(exc)
        return state


# ---------------------------------------------------------------------------
# Generation actions
# ---------------------------------------------------------------------------


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("generate-drs", "n_clicks"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def generate_drs_action(_n_clicks, session_data, state):
    # Dash can fire callbacks when interactive components are dynamically
    # inserted into the layout. A real button click always has n_clicks >= 1.
    if not _n_clicks:
        raise PreventUpdate
    state = dict(state or fresh_state())
    brd = state.get("uploaded_brd")
    if not brd or state.get("ingestion_status") != "done":
        raise PreventUpdate
    try:
        _start_background_job(
            session_data["session_uid"],
            "drs",
            _backend_callable("src.fetch_drs_pipeline", "generate_drs"),
            {"brs_name": Path(brd).stem},
        )
        state["drs_status"] = "running"
        state["drs_error"] = None
    except Exception as exc:
        state["drs_status"] = "error"
        state["drs_error"] = str(exc)
    return state


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("generate-jira", "n_clicks"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def generate_jira_action(_n_clicks, session_data, state):
    # Ignore the initial component-insertion event; only a real click starts Jira.
    if not _n_clicks:
        raise PreventUpdate
    state = dict(state or fresh_state())
    brd = state.get("uploaded_brd")
    if not brd or state.get("ingestion_status") != "done":
        raise PreventUpdate
    try:
        _start_background_job(
            session_data["session_uid"],
            "jira",
            _backend_callable("src.fetch_jira_pipeline", "generate_jira_document_pipeline"),
            {"brs_name": Path(brd).stem},
        )
        state["jira_status"] = "running"
        state["jira_error"] = None
    except Exception as exc:
        state["jira_status"] = "error"
        state["jira_error"] = str(exc)
    return state


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("generate-estimation", "n_clicks"),
    State("qa-support", "value"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def generate_estimation_action(_n_clicks, qa_support_values, session_data, state):
    # Ignore the initial component-insertion event; only a real click starts estimation.
    if not _n_clicks:
        raise PreventUpdate
    state = dict(state or fresh_state())
    brd = state.get("uploaded_brd")
    if not brd or state.get("ingestion_status") != "done":
        raise PreventUpdate
    if any(state.get(f"{name}_status") == "running" for name in ("ingestion", "drs", "jira")):
        raise PreventUpdate
    try:
        _start_background_job(
            session_data["session_uid"],
            "estimation",
            _backend_callable("estimation_generation.pipeline", "generate_estimation"),
            {
                "brs_name": Path(brd).stem,
                "qa_support": bool(qa_support_values and "included" in qa_support_values),
            },
        )
        state["estimation_status"] = "running"
        state["estimation_error"] = None
    except Exception as exc:
        state["estimation_status"] = "error"
        state["estimation_error"] = str(exc)
    return state


# ---------------------------------------------------------------------------
# Chatbot
# ---------------------------------------------------------------------------


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("chat-send", "n_clicks"),
    Input("chat-input", "n_submit"),
    State("chat-input", "value"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def start_chat(_send_clicks, _submit_count, question, session_data, state):
    if not question or not question.strip():
        raise PreventUpdate
    state = dict(state or fresh_state())
    if state.get("ingestion_status") != "done":
        raise PreventUpdate
    if state.get("chat_job"):
        raise PreventUpdate
    if any(state.get(f"{name}_status") == "running" for name in ("ingestion", "drs", "jira", "estimation")):
        raise PreventUpdate

    session_uid = session_data["session_uid"]
    question = question.strip()
    chat_history = list(state.get("chat_history", []))
    prior_history = [m for m in chat_history if not m.get("pending")]
    message_id = uuid.uuid4().hex
    chat_history.append({"role": "user", "content": question})
    chat_history.append({"role": "assistant", "content": "", "pending": True, "message_id": message_id})

    try:
        kb_path = state.get("chat_kb_path")
        job_name = f"chat:{message_id}"

        def _chat_target():
            executor_cls = _backend_callable("chatbot_feature.chatbot_executor", "ChatbotExecutor")
            executor = executor_cls(
                knowledge_base_path=kb_path,
                candidate_top_k=10,
                rerank_top_k=3,
            )
            return executor.answer(question=question, chat_history=prior_history)

        _start_background_job(session_uid, job_name, _chat_target, {})
        state["chat_history"] = chat_history
        state["chat_job"] = message_id
    except Exception as exc:
        chat_history[-1] = {
            "role": "assistant",
            "content": f"I could not answer the question because: {exc}",
            "pending": False,
            "message_id": message_id,
        }
        state["chat_history"] = chat_history
        state["chat_job"] = None
    return state


# ---------------------------------------------------------------------------
# Background polling / persistence fallback
# ---------------------------------------------------------------------------


@app.callback(
    Output("app-state", "data", allow_duplicate=True),
    Input("job-poller", "n_intervals"),
    State("session-store", "data"),
    State("app-state", "data"),
    prevent_initial_call=True,
)
def poll_jobs(_n_intervals, session_data, state):
    state = dict(state or fresh_state())
    session_uid = session_data["session_uid"]
    changed = False

    for job_name, payload in _drain_session_results(session_uid):
        status = payload.get("status")
        result = payload.get("result")
        error = payload.get("error")

        if job_name == "ingestion":
            state["ingestion_status"] = status
            state["ingestion_error"] = error
            if status == "done":
                state["chat_kb_path"] = _extract_kb_path(result)
            changed = True
        elif job_name == "drs":
            state["drs_status"] = status
            state["drs_error"] = error
            if status == "done":
                word = _result_path(result)
                if word:
                    state["drs_word"] = str(word)
                    md = _result_markdown(result)
                    if not md:
                        md = word.parent / f"{word.stem}.md"
                    if not md.is_file():
                        md = _docx_to_markdown(word, md)
                    state["drs_md"] = str(md)
            changed = True
        elif job_name == "jira":
            state["jira_status"] = status
            state["jira_error"] = error
            if status == "done":
                word = _result_path(result)
                if word:
                    state["jira_word"] = str(word)
                    md = _result_markdown(result)
                    if not md:
                        md = word.parent / f"{word.stem}.md"
                    if not md.is_file():
                        md = _docx_to_markdown(word, md)
                    state["jira_md"] = str(md)
            changed = True
        elif job_name == "estimation":
            state["estimation_status"] = status
            state["estimation_error"] = error
            if status == "done":
                workbook = _result_path(result)
                if workbook:
                    state["estimation_xlsx"] = str(workbook)
            changed = True
        elif job_name.startswith("chat:"):
            message_id = job_name.split(":", 1)[1]
            history = list(state.get("chat_history", []))
            for message in history:
                if message.get("message_id") == message_id and message.get("pending"):
                    if status == "done":
                        message["content"] = str(result)
                    else:
                        message["content"] = f"I could not answer the question because: {error or 'Unknown error'}"
                    message["pending"] = False
                    break
            state["chat_history"] = history
            if state.get("chat_job") == message_id:
                state["chat_job"] = None
            changed = True

    # Recovery signal mirrors the Streamlit persisted-artifact fallback.
    uploaded = state.get("uploaded_brd")
    if uploaded:
        folder = OUTPUT_DIR / Path(uploaded).stem
        if folder.is_dir():
            persisted = {
                "drs": folder / "generated_drs.docx",
                "jira": folder / "generated_jira_user_stories.docx",
                "estimation": _find_estimation_artifact(folder),
            }
            for name, path in persisted.items():
                if path and path.is_file() and state.get(f"{name}_status") == "running":
                    state[f"{name}_status"] = "done"
                    if name == "drs":
                        state["drs_word"] = str(path)
                    elif name == "jira":
                        state["jira_word"] = str(path)
                    elif name == "estimation":
                        state["estimation_xlsx"] = str(path)
                    changed = True

    if not changed:
        raise PreventUpdate
    return state


# ---------------------------------------------------------------------------
# Downloads
# ---------------------------------------------------------------------------


def _send_path(path: Optional[str]):
    safe = _safe_existing_path(path or "")
    if not safe:
        raise PreventUpdate
    return dcc.send_file(str(safe))


@app.callback(Output("file-download", "data"),
              Input("download-drs-word", "n_clicks"),
              Input("download-drs-md", "n_clicks"),
              Input("download-jira-word", "n_clicks"),
              Input("download-jira-md", "n_clicks"),
              State("app-state", "data"),
              prevent_initial_call=True)
def download_artifact(*args):
    # Dynamic layout insertion can cause an input callback to be evaluated
    # without an actual click. Never treat None/0 as a download request.
    click_values = args[:-1]
    if not any(isinstance(value, (int, float)) and value > 0 for value in click_values):
        raise PreventUpdate

    state = args[-1] or fresh_state()
    mapping = {
        "download-drs-word": "drs_word",
        "download-drs-md": "drs_md",
        "download-jira-word": "jira_word",
        "download-jira-md": "jira_md",
    }
    key = ctx.triggered_id
    return _send_path(state.get(mapping.get(key, "")))


@app.callback(
    Output("file-download", "data", allow_duplicate=True),
    Input({"type": "estimation-download", "path": ALL}, "n_clicks"),
    prevent_initial_call=True,
)
def download_estimation(_clicks):
    """Download the estimation from the right-side ready link."""
    if not isinstance(_clicks, list) or not any(
        isinstance(value, (int, float)) and value > 0 for value in _clicks
    ):
        raise PreventUpdate
    triggered = ctx.triggered_id
    if not isinstance(triggered, dict):
        raise PreventUpdate
    return _send_path(triggered.get("path"))


@app.callback(
    Output("file-download", "data", allow_duplicate=True),
    Input({"type": "dashboard-download", "kind": ALL, "path": ALL}, "n_clicks"),
    prevent_initial_call=True,
)
def download_dashboard(_clicks):
    # The dashboard download buttons are created dynamically. Their callback
    # must only respond to an actual positive click, never to component
    # insertion/re-rendering.
    if not isinstance(_clicks, list) or not any(
        isinstance(value, (int, float)) and value > 0 for value in _clicks
    ):
        raise PreventUpdate

    triggered = ctx.triggered_id
    if not isinstance(triggered, dict):
        raise PreventUpdate
    return _send_path(triggered.get("path"))


# ---------------------------------------------------------------------------
# Static Flask health endpoint
# ---------------------------------------------------------------------------


@app.server.get("/health")
def health():
    return {"status": "ok"}


def open_browser():
    webbrowser.open_new("http://127.0.0.1:8050/")
if __name__ == "__main__":

    Timer(1,open_browser).start() 

    host = os.getenv("DASH_HOST", "127.0.0.1")
    port = int(os.getenv("DASH_PORT", "8050"))
    debug = os.getenv("DASH_DEBUG", "false").lower() == "true"
    app.run(host=host, port=port, debug=debug)
