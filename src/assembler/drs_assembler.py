from __future__ import annotations

import json
import logging
import os
from copy import deepcopy
from typing import Any, Dict, List, Optional, Union

from docx import Document
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph


# =============================================================================
# PROJECT PATHS
# =============================================================================

BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.dirname(
            os.path.abspath(__file__)
        )
    )
)

TEMPLATE_PATH = os.path.join(
    BASE_DIR,
    "data",
    "uploads",
    "DRS_Template.docx",
)

ANSWERS_PATH = os.path.join(
    BASE_DIR,
    "query_results",
    "all_query_results.json",
)

OUTPUT_DIR = os.path.join(
    BASE_DIR,
    "output",
)

OUTPUT_PATH = os.path.join(
    OUTPUT_DIR,
    "generated_drs.docx",
)

LOG_DIR = os.path.join(
    BASE_DIR,
    "logs",
)

LOG_FILE = os.path.join(
    LOG_DIR,
    "drs_assembly.log",
)


# =============================================================================
# LOGGING
# =============================================================================

def configure_logging() -> logging.Logger:
    """Configure centralized DRS assembly logging."""

    os.makedirs(
        LOG_DIR,
        exist_ok=True,
    )

    logger = logging.getLogger("drs_assembler")

    logger.setLevel(logging.INFO)

    # Avoid duplicate handlers on repeated imports/runs.
    logger.handlers.clear()

    formatter = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(name)s | %(message)s"
    )

    file_handler = logging.FileHandler(
        LOG_FILE,
        mode="a",
        encoding="utf-8",
    )

    console_handler = logging.StreamHandler()

    file_handler.setFormatter(formatter)
    console_handler.setFormatter(formatter)

    logger.addHandler(file_handler)
    logger.addHandler(console_handler)

    return logger


logger = logging.getLogger("drs_assembler")


# =============================================================================
# TYPES
# =============================================================================

QueryResults = Dict[str, Dict[str, Any]]


# =============================================================================
# LOAD QUERY RESULTS
# =============================================================================

def load_query_results(
    path: str,
) -> QueryResults:
    """Load all generated query results."""

    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"Query results file not found: {path}"
        )

    with open(
        path,
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    if not isinstance(data, dict):
        raise ValueError(
            "all_query_results.json must contain a JSON object."
        )

    return data


def get_answer(
    results: QueryResults,
    query_id: str,
) -> Any:
    """Return the answer associated with a query ID."""

    result = results.get(query_id)

    if not result:
        return None

    return result.get("answer")


# =============================================================================
# GENERAL TEXT HELPERS
# =============================================================================

def stringify(
    value: Any,
) -> str:
    """Convert a value into text suitable for a DOCX cell/paragraph."""

    if value is None:
        return ""

    if isinstance(value, str):
        return value

    if isinstance(value, list):
        return ", ".join(
            stringify(item)
            for item in value
        )

    if isinstance(value, dict):
        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
        )

    return str(value)


def replace_paragraph_text(
    paragraph: Paragraph,
    text: Any,
) -> None:
    """
    Replace paragraph contents while preserving the paragraph itself.

    Values coming from LLM query results may be None or another non-string
    type, so they are normalized before being written to python-docx.
    The first existing run is reused so paragraph-level formatting remains
    attached to the paragraph.
    """

    text_value = stringify(text)

    if paragraph.runs:
        paragraph.runs[0].text = text_value

        for run in paragraph.runs[1:]:
            run.text = ""

    else:
        paragraph.add_run(text_value)


# =============================================================================
# TABLE HELPERS
# =============================================================================

def normalize_cell_text(
    cell: _Cell,
) -> str:
    """Normalize a table cell's text for header matching."""

    return " ".join(
        cell.text.split()
    ).strip().lower()


def get_table_headers(
    table: Table,
) -> List[str]:
    """Return normalized headers from the first row."""

    if not table.rows:
        return []

    return [
        normalize_cell_text(cell)
        for cell in table.rows[0].cells
    ]


def find_table(
    document: Document,
    required_headers: List[str],
) -> Optional[Table]:
    """
    Find a table by its header names.

    This deliberately does not rely on table position/index.
    """

    required = {
        header.strip().lower()
        for header in required_headers
    }

    for table in document.tables:
        headers = set(
            get_table_headers(table)
        )

        if required.issubset(headers):
            return table

    return None


def get_template_data_row(
    table: Table,
) -> Optional[Any]:
    """
    Return a deep copy of the first template data row.

    The first row is assumed to be the header.
    The second row, when present, is the placeholder/template row whose
    formatting should be reused for generated rows.
    """

    if len(table.rows) < 2:
        return None

    return deepcopy(
        table.rows[1]._tr
    )


def clear_data_rows(
    table: Table,
) -> None:
    """Remove all rows except the header row."""

    while len(table.rows) > 1:
        table._tbl.remove(
            table.rows[-1]._tr
        )


def append_template_row(
    table: Table,
    template_row: Any,
) -> List[_Cell]:
    """
    Append a fully cloned template row so formatting is preserved.
    """

    cloned_row = deepcopy(
        template_row
    )

    table._tbl.append(
        cloned_row
    )

    return list(
        table.rows[-1].cells
    )


def write_cell(
    cell: _Cell,
    value: Any,
) -> None:
    """
    Replace the text in a cell while preserving the cell's formatting.
    """

    text = stringify(value)

    if not cell.paragraphs:
        cell.text = text
        return

    paragraph = cell.paragraphs[0]

    if paragraph.runs:
        paragraph.runs[0].text = text

        for run in paragraph.runs[1:]:
            run.text = ""

    else:
        paragraph.add_run(text)

    # Blank additional paragraphs instead of deleting them, preserving
    # the original cell structure.
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def populate_table_rows(
    table: Table,
    rows: List[List[Any]],
) -> None:
    """
    Populate a table using the formatting of its original placeholder row.

    Header row is preserved.
    """

    template_row = get_template_data_row(
        table
    )

    if template_row is None:
        raise ValueError(
            "Table does not contain a template data row."
        )

    clear_data_rows(
        table
    )

    # If the answer is an empty list, preserve one blank template row so
    # the original table structure remains visible.
    rows_to_add = rows or [
        [""] * len(
            table.rows[0].cells
        )
    ]

    for row_values in rows_to_add:
        cells = append_template_row(
            table,
            template_row,
        )

        for index, value in enumerate(
            row_values
        ):
            if index >= len(cells):
                break

            write_cell(
                cells[index],
                value,
            )


# =============================================================================
# SIMPLE TEXT SECTIONS
# =============================================================================

def populate_project_metadata(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate project name and version."""

    metadata = get_answer(
        results,
        "Q001",
    )

    if not isinstance(
        metadata,
        dict,
    ):
        logger.warning(
            "Q001 answer is not an object."
        )
        return

    project_name = metadata.get(
        "project_name",
        "",
    )

    version = metadata.get(
        "version",
        "",
    )

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text == "Project:":
            replace_paragraph_text(
                paragraph,
                f"Project: {project_name}",
            )

        elif text == "Version:":
            replace_paragraph_text(
                paragraph,
                f"Version: {version}",
            )


def populate_executive_summary(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate Executive Summary."""

    data = get_answer(
        results,
        "Q002",
    )

    if not isinstance(
        data,
        dict,
    ):
        logger.warning(
            "Q002 answer is not an object."
        )
        return

    summary = data.get(
        "summary",
        "",
    )

    for paragraph in document.paragraphs:

        if paragraph.text.strip() == "[Enter executive summary]":
            replace_paragraph_text(
                paragraph,
                summary,
            )
            return


def populate_business_background(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate Business Background."""

    data = get_answer(
        results,
        "Q003",
    )

    if not isinstance(
        data,
        dict,
    ):
        logger.warning(
            "Q003 answer is not an object."
        )
        return

    background = data.get(
        "background",
        "",
    )

    for paragraph in document.paragraphs:

        if paragraph.text.strip() == "[Enter business background]":
            replace_paragraph_text(
                paragraph,
                background,
            )
            return


def populate_business_objectives(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate Business Objectives."""

    objectives = get_answer(
        results,
        "Q004",
    )

    if not isinstance(
        objectives,
        list,
    ):
        logger.warning(
            "Q004 answer is not an array."
        )
        return

    lines: List[str] = []

    for item in objectives:

        if not isinstance(
            item,
            dict,
        ):
            continue

        objective_id = item.get(
            "objective_id",
        )

        objective = item.get(
            "objective",
            "",
        )

        if objective_id:
            lines.append(
                f"{objective_id} {objective}"
            )
        else:
            lines.append(
                objective
            )

    for paragraph in document.paragraphs:

        if paragraph.text.strip() == "[Enter business objectives]":
            replace_paragraph_text(
                paragraph,
                "\n".join(lines),
            )
            return


def populate_scope(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate the separate in-scope and out-of-scope placeholders."""

    scope = get_answer(
        results,
        "Q005",
    )

    out_of_scope = get_answer(
        results,
        "Q006",
    )

    if not isinstance(
        scope,
        dict,
    ):
        logger.warning(
            "Q005 answer is not an object."
        )
        return

    in_scope = scope.get(
        "in_scope",
        [],
    )

    if not isinstance(
        in_scope,
        list,
    ):
        in_scope = [
            in_scope
        ]

    out_scope: List[str] = []

    if isinstance(
        out_of_scope,
        list,
    ):
        for item in out_of_scope:

            if isinstance(
                item,
                dict,
            ):
                value = item.get(
                    "item",
                    "",
                )

                if value:
                    out_scope.append(
                        str(value)
                    )

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text == "In Scope: [Enter in-scope items]":
            replace_paragraph_text(
                paragraph,
                (
                    "In Scope: "
                    + ", ".join(
                        str(item)
                        for item in in_scope
                    )
                ),
            )

        elif text == "Out of Scope: [Enter out-of-scope items]":
            replace_paragraph_text(
                paragraph,
                (
                    "Out of Scope: "
                    + ", ".join(out_scope)
                ),
            )


def populate_stakeholders(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate stakeholder paragraph."""

    stakeholders = get_answer(
        results,
        "Q007",
    )

    if not isinstance(
        stakeholders,
        list,
    ):
        logger.warning(
            "Q007 answer is not an array."
        )
        return

    lines: List[str] = []

    for item in stakeholders:

        if not isinstance(
            item,
            dict,
        ):
            continue

        name = item.get(
            "name_or_role",
            "",
        )

        responsibility = item.get(
            "responsibility",
            "",
        )

        if responsibility:
            lines.append(
                f"{name} – {responsibility}"
            )
        else:
            lines.append(
                str(name)
            )

    for paragraph in document.paragraphs:

        if paragraph.text.strip() == "[Enter stakeholders and roles]":
            replace_paragraph_text(
                paragraph,
                "\n".join(lines),
            )
            return


def populate_review_recommendations(
    document: Document,
    results: QueryResults,
) -> None:
    """Populate the Review Recommendations placeholder."""

    recommendations = get_answer(
        results,
        "Q025",
    )

    if not isinstance(
        recommendations,
        list,
    ):
        logger.warning(
            "Q025 answer is not an array or is unavailable."
        )
        return

    lines: List[str] = []

    for item in recommendations:

        if not isinstance(
            item,
            dict,
        ):
            continue

        recommendation = item.get(
            "recommendation",
            "",
        )

        if recommendation:
            lines.append(
                str(recommendation)
            )

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text.startswith(
            "[Enter review recommendations"
        ):
            replace_paragraph_text(
                paragraph,
                "\n".join(lines),
            )
            logger.info(
                "Review Recommendations populated."
            )
            return

    logger.warning(
        "Review Recommendations placeholder not found."
    )


# =============================================================================
# EPICS TABLE
# =============================================================================

def populate_epics_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Epics table from Q018."""

    epics = get_answer(
        results,
        "Q018",
    )

    if epics is None:
        logger.warning(
            "Q018 has no answer."
        )
        return

    # Some LLM responses incorrectly return one object rather than an array.
    # The assembler accepts both so the document can still be generated.
    if isinstance(
        epics,
        dict,
    ):
        epics = [
            epics
        ]

    if not isinstance(
        epics,
        list,
    ):
        logger.warning(
            "Q018 answer has unsupported structure."
        )
        return

    rows: List[List[Any]] = []

    for epic in epics:

        if not isinstance(
            epic,
            dict,
        ):
            continue

        epic_id = epic.get(
            "epic_id",
            "",
        )

        epic_name = epic.get(
            "epic_name",
            "",
        )

        description = epic.get(
            "description",
            "",
        )

        linked_fr_ids = stringify(
            epic.get(
                "linked_fr_ids",
                [],
            )
        )

        name_description = str(
            epic_name
        )

        if description:
            name_description += (
                f"\n{description}"
            )

        if linked_fr_ids:
            name_description += (
                f"\nFR: {linked_fr_ids}"
            )

        rows.append(
            [
                epic_id,
                name_description,
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# FUNCTIONAL REQUIREMENTS TABLE
# =============================================================================

def populate_functional_requirements_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Functional Requirements from Q015."""

    requirements = get_answer(
        results,
        "Q015",
    )

    if not isinstance(
        requirements,
        list,
    ):
        logger.warning(
            "Q015 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for requirement in requirements:

        if not isinstance(
            requirement,
            dict,
        ):
            continue

        rows.append(
            [
                requirement.get(
                    "fr_id",
                    "",
                ),
                requirement.get(
                    "title",
                    "",
                ),
                requirement.get(
                    "actor",
                    "",
                ),
                requirement.get(
                    "priority",
                    "",
                ),
                requirement.get(
                    "source_reference",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# NFR TABLE
# =============================================================================

def populate_nfr_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate NFR table from Q016."""

    requirements = get_answer(
        results,
        "Q016",
    )

    if not isinstance(
        requirements,
        list,
    ):
        logger.warning(
            "Q016 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for requirement in requirements:

        if not isinstance(
            requirement,
            dict,
        ):
            continue

        rows.append(
            [
                requirement.get(
                    "nfr_id",
                    "",
                ),
                requirement.get(
                    "category",
                    "",
                ),
                requirement.get(
                    "requirement",
                    "",
                ),
                requirement.get(
                    "priority",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# BUSINESS RULES TABLE
# =============================================================================

def populate_business_rules_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Business Rules from Q017."""

    rules = get_answer(
        results,
        "Q017",
    )

    if not isinstance(
        rules,
        list,
    ):
        logger.warning(
            "Q017 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for rule in rules:

        if not isinstance(
            rule,
            dict,
        ):
            continue

        rows.append(
            [
                rule.get(
                    "rule_id",
                    "",
                ),
                rule.get(
                    "rule_statement",
                    "",
                ),
                rule.get(
                    "source_reference",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# ACCEPTANCE CRITERIA TABLE
# =============================================================================

def populate_acceptance_criteria_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Acceptance Criteria from Q020."""

    criteria = get_answer(
        results,
        "Q020",
    )

    if not isinstance(
        criteria,
        list,
    ):
        logger.warning(
            "Q020 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for criterion in criteria:

        if not isinstance(
            criterion,
            dict,
        ):
            continue

        rows.append(
            [
                criterion.get(
                    "ac_id",
                    "",
                ),
                criterion.get(
                    "given",
                    "",
                ),
                criterion.get(
                    "when",
                    "",
                ),
                criterion.get(
                    "then",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# TRACEABILITY TABLE
# =============================================================================

def populate_traceability_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Traceability Matrix from Q024."""

    matrix = get_answer(
        results,
        "Q024",
    )

    if not isinstance(
        matrix,
        list,
    ):
        logger.warning(
            "Q024 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for item in matrix:

        if not isinstance(
            item,
            dict,
        ):
            continue

        rows.append(
            [
                item.get(
                    "business_objective",
                    "",
                ),
                item.get(
                    "functional_requirement",
                    "",
                ),
                item.get(
                    "user_story",
                    "",
                ),
                item.get(
                    "acceptance_criteria",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# CLARIFICATION LOG TABLE
# =============================================================================

def populate_clarification_table(
    table: Table,
    results: QueryResults,
) -> None:
    """Populate Clarification Log from Q023."""

    clarifications = get_answer(
        results,
        "Q023",
    )

    if not isinstance(
        clarifications,
        list,
    ):
        logger.warning(
            "Q023 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for item in clarifications:

        if not isinstance(
            item,
            dict,
        ):
            continue

        rows.append(
            [
                item.get(
                    "cq_id",
                    "",
                ),
                item.get(
                    "question",
                    "",
                ),
                item.get(
                    "reason",
                    "",
                ),
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# USER STORIES TABLE
# =============================================================================

def populate_user_stories_table(
    table: Table,
    results: QueryResults,
) -> None:
    """
    Populate the dedicated User Stories table from Q019.

    Template structure:
        User Story ID | User Story | Related Epic / FR
    """

    stories = get_answer(
        results,
        "Q019",
    )

    if not isinstance(
        stories,
        list,
    ):
        logger.warning(
            "Q019 answer is not an array."
        )
        return

    rows: List[List[Any]] = []

    for story in stories:

        if not isinstance(
            story,
            dict,
        ):
            continue

        story_id = story.get(
            "story_id",
            "",
        )

        role = story.get(
            "role",
            "",
        )

        capability = story.get(
            "capability",
            "",
        )

        benefit = story.get(
            "benefit",
            "",
        )

        linked_fr_ids = stringify(
            story.get(
                "linked_fr_ids",
                [],
            )
        )

        user_story_text = (
            f"As a {role}, "
            f"I want to {capability}, "
            f"so that {benefit}."
        )

        related_text = (
            f"FR: {linked_fr_ids}"
            if linked_fr_ids
            else ""
        )

        rows.append(
            [
                story_id,
                user_story_text,
                related_text,
            ]
        )

    populate_table_rows(
        table,
        rows,
    )


# =============================================================================
# TABLE DISCOVERY + POPULATION
# =============================================================================

def populate_all_tables(
    document: Document,
    results: QueryResults,
) -> None:
    """
    Discover each DRS table from its headers and populate it.

    Table order is deliberately ignored.
    """

    mappings = [
        (
            "Epics",
            [
                "epic id",
                "epic name / description",
            ],
            populate_epics_table,
        ),
        (
            "Functional Requirements",
            [
                "fr id",
                "title",
                "actor",
                "priority",
                "source",
            ],
            populate_functional_requirements_table,
        ),
        (
            "Non Functional Requirements",
            [
                "nfr id",
                "category",
                "requirement",
                "priority",
            ],
            populate_nfr_table,
        ),
        (
            "Business Rules",
            [
                "rule id",
                "rule statement",
                "source",
            ],
            populate_business_rules_table,
        ),
        (
            "Acceptance Criteria",
            [
                "ac id",
                "given",
                "when",
                "then",
            ],
            populate_acceptance_criteria_table,
        ),
        (
            "Traceability Matrix",
            [
                "business objective",
                "fr",
                "user story",
                "ac",
            ],
            populate_traceability_table,
        ),
        (
            "Clarification Log",
            [
                "cq id",
                "question",
                "reason",
            ],
            populate_clarification_table,
        ),
        (
            "User Stories",
            [
                "user story id",
                "user story",
                "related epic / fr",
            ],
            populate_user_stories_table,
        ),
    ]

    for section_name, headers, handler in mappings:

        table = find_table(
            document,
            headers,
        )

        if table is None:
            logger.warning(
                "Could not find table for section: %s",
                section_name,
            )
            continue

        logger.info(
            "Populating table: %s",
            section_name,
        )

        handler(
            table,
            results,
        )


# =============================================================================
# MAIN
# =============================================================================

def generate_drs(output_dir: Optional[Union[str, os.PathLike[str]]] = None) -> str:
    """Generate a completed DRS using the fixed template and answer paths."""

    configure_logging()

    output_directory = (
        os.fspath(output_dir)
        if output_dir is not None
        else OUTPUT_DIR
    )
    output_path = os.path.join(
        output_directory,
        "generated_drs.docx",
    )

    logger.info("=" * 80)
    logger.info("DRS ASSEMBLY STARTED")
    logger.info("=" * 80)

    logger.info(
        "Template: %s",
        TEMPLATE_PATH,
    )

    logger.info(
        "Answers: %s",
        ANSWERS_PATH,
    )

    logger.info(
        "Output: %s",
        output_path,
    )

    # -------------------------------------------------------------------------
    # Validate inputs
    # -------------------------------------------------------------------------

    if not os.path.isfile(
        TEMPLATE_PATH
    ):
        logger.error(
            "DRS template not found: %s",
            TEMPLATE_PATH,
        )

        raise FileNotFoundError(
            f"DRS template not found: {TEMPLATE_PATH}"
        )

    if not os.path.isfile(
        ANSWERS_PATH
    ):
        logger.error(
            "Query results not found: %s",
            ANSWERS_PATH,
        )

        raise FileNotFoundError(
            f"Query results not found: {ANSWERS_PATH}"
        )

    # -------------------------------------------------------------------------
    # Load query results
    # -------------------------------------------------------------------------

    results = load_query_results(
        ANSWERS_PATH
    )

    logger.info(
        "Loaded %d query results.",
        len(results),
    )

    # -------------------------------------------------------------------------
    # Load template
    # -------------------------------------------------------------------------

    document = Document(
        TEMPLATE_PATH
    )

    logger.info(
        "Loaded DRS template: paragraphs=%d, tables=%d",
        len(document.paragraphs),
        len(document.tables),
    )

    # -------------------------------------------------------------------------
    # Populate paragraph-based sections
    # -------------------------------------------------------------------------

    logger.info(
        "Populating Project Metadata..."
    )

    populate_project_metadata(
        document,
        results,
    )

    logger.info(
        "Populating Executive Summary..."
    )

    populate_executive_summary(
        document,
        results,
    )

    logger.info(
        "Populating Business Background..."
    )

    populate_business_background(
        document,
        results,
    )

    logger.info(
        "Populating Business Objectives..."
    )

    populate_business_objectives(
        document,
        results,
    )

    logger.info(
        "Populating Scope..."
    )

    populate_scope(
        document,
        results,
    )

    logger.info(
        "Populating Stakeholders..."
    )

    populate_stakeholders(
        document,
        results,
    )

    logger.info(
        "Populating Review Recommendations..."
    )

    populate_review_recommendations(
        document,
        results,
    )

    # -------------------------------------------------------------------------
    # Populate tables
    # -------------------------------------------------------------------------

    logger.info(
        "Populating DRS tables..."
    )

    populate_all_tables(
        document,
        results,
    )

    # -------------------------------------------------------------------------
    # Save output
    # -------------------------------------------------------------------------

    os.makedirs(
        output_directory,
        exist_ok=True,
    )

    logger.info(
        "Saving generated DRS..."
    )

    document.save(
        output_path
    )

    logger.info(
        "DRS generated successfully: %s",
        output_path,
    )

    logger.info("=" * 80)
    logger.info("DRS ASSEMBLY COMPLETED")
    logger.info("=" * 80)

    print()
    print("=" * 70)
    print("DRS generated successfully")
    print("=" * 70)
    print(
        f"Output: {output_path}"
    )
    print(
        f"Log:    {LOG_FILE}"
    )
    print("=" * 70)

    return output_path


if __name__ == "__main__":
    generate_drs()