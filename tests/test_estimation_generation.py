from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from estimation_generation.document_builder import write_estimation_report
from estimation_generation.scoring import FACTORS, FactorScore, calculate_weighted_score


class EstimationGenerationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.factors = [
            FactorScore(key, label, weight, 3, f"{label} is supported by the extracted BRD evidence.", ["Example source-grounded observation."])
            for key, label, weight in FACTORS
        ]

    def test_weighted_score_uses_configured_weights(self) -> None:
        self.assertEqual(calculate_weighted_score(self.factors), 3.0)

    def test_report_writes_markdown_and_docx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            docx_path, markdown_path = write_estimation_report(Path(directory), "sample_brs", self.factors, 3.0)
            self.assertTrue(docx_path.is_file())
            self.assertTrue(markdown_path.is_file())
            self.assertIn("Final Complexity Score", markdown_path.read_text(encoding="utf-8"))
            document = Document(docx_path)
            self.assertIn("BRD Delivery Estimation", document.paragraphs[0].text)
            self.assertEqual(len(document.tables), 1)


if __name__ == "__main__":
    unittest.main()
