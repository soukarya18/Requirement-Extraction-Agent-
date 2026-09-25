from __future__ import annotations

import os
import re
from typing import Any, Dict, List, Sequence

import fitz
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from modules.config import CHUNK_OVERLAP, CHUNK_SIZE, get_logger


logger = get_logger("document_loader")


# ---------------------------------------------------------------------------
# Heading detection
# ---------------------------------------------------------------------------

_NUMBERED_HEADING_RE = re.compile(
    r"^\d+(?:\.\d+)*\.?\s+.+$"
)


def _is_numbered_heading(text: str) -> bool:
    """Detect numbered section/subsection headings."""
    text = " ".join(text.strip().split())
    return bool(
        _NUMBERED_HEADING_RE.match(text)
        and len(text) <= 180
    )


def _normalize_heading(text: str) -> str:
    """Normalize heading text for metadata and chunk prefixes."""
    return " ".join(text.strip().split())


def _docx_heading_level(paragraph: Paragraph) -> int:
    """Return the best-effort DOCX heading level."""
    try:
        style_name = (paragraph.style.name or "").strip()
        match = re.search(r"heading\s*(\d+)", style_name, re.I)
        if match:
            return int(match.group(1))
    except Exception:
        pass

    text = paragraph.text.strip()
    if _is_numbered_heading(text):
        marker = text.split(maxsplit=1)[0].rstrip(".")
        return max(1, marker.count(".") + 1)

    return 1


def _is_docx_heading(paragraph: Paragraph) -> bool:
    """Detect a DOCX heading using Word structure plus safe fallbacks."""
    text = paragraph.text.strip()
    if not text:
        return False

    try:
        style_name = (paragraph.style.name or "").lower()
        if style_name.startswith("heading"):
            return True
    except Exception:
        pass

    if _is_numbered_heading(text):
        return True

    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return False

    bold = all(bool(run.bold) for run in runs)
    sizes = [
        float(run.font.size.pt)
        for run in runs
        if run.font.size is not None
    ]

    return bool(
        bold
        and sizes
        and max(sizes) >= 14.0
        and len(text) <= 180
    )


def _iter_docx_blocks(doc: Any) -> Sequence[Any]:
    """Yield DOCX paragraphs and tables in original document order."""
    paragraph_map = {p._p: p for p in doc.paragraphs}
    table_map = {t._tbl: t for t in doc.tables}

    blocks: List[Any] = []
    for element in doc.element.body.iterchildren():
        if element in paragraph_map:
            blocks.append(paragraph_map[element])
        elif element in table_map:
            blocks.append(table_map[element])
    return blocks


# ---------------------------------------------------------------------------
# DOCX structural extraction
# ---------------------------------------------------------------------------


def _load_docx_sections(path: str) -> List[Dict[str, Any]]:
    """Extract DOCX content in document order, grouped by section heading."""
    doc = DocxDocument(path)
    blocks: List[Dict[str, Any]] = []

    current_section = "Document"
    current_level = 0

    for block in _iter_docx_blocks(doc):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue

            if _is_docx_heading(block):
                current_section = _normalize_heading(text)
                current_level = _docx_heading_level(block)
                blocks.append(
                    {
                        "section": current_section,
                        "section_level": current_level,
                        "content": current_section,
                    }
                )
            else:
                blocks.append(
                    {
                        "section": current_section,
                        "section_level": current_level,
                        "content": text,
                    }
                )

        elif isinstance(block, Table):
            rows: List[str] = []
            for row in block.rows:
                cells = [
                    " ".join(cell.text.split())
                    for cell in row.cells
                ]
                row_text = " | ".join(cells)
                if row_text.strip():
                    rows.append(row_text)

            if rows:
                blocks.append(
                    {
                        "section": current_section,
                        "section_level": current_level,
                        "content": "\n".join(rows),
                    }
                )

    return blocks


# ---------------------------------------------------------------------------
# PDF structural extraction
# ---------------------------------------------------------------------------


def _extract_pdf_lines(path: str) -> List[Dict[str, Any]]:
    """Extract PDF lines with font and page information."""
    pdf = fitz.open(path)
    lines: List[Dict[str, Any]] = []

    try:
        for page_number, page in enumerate(pdf, start=1):
            page_dict = page.get_text("dict")

            for block in page_dict.get("blocks", []):
                if block.get("type") != 0:
                    continue

                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    text = " ".join(
                        span.get("text", "").strip()
                        for span in spans
                        if span.get("text", "").strip()
                    ).strip()

                    if not text:
                        continue

                    sizes = [
                        float(span.get("size", 0.0))
                        for span in spans
                        if span.get("text", "").strip()
                    ]

                    fonts = [
                        str(span.get("font", "")).lower()
                        for span in spans
                        if span.get("text", "").strip()
                    ]

                    lines.append(
                        {
                            "page": page_number,
                            "text": text,
                            "font_size": max(sizes) if sizes else 0.0,
                            "bold": bool(fonts) and all("bold" in f for f in fonts),
                        }
                    )
    finally:
        pdf.close()

    return lines


def _is_pdf_heading(
    line: Dict[str, Any],
    median_font_size: float,
) -> bool:
    """Detect a likely PDF heading from numbering/font structure."""
    text = str(line["text"]).strip()
    size = float(line["font_size"])
    bold = bool(line["bold"])

    if not text or len(text) > 180:
        return False

    if _is_numbered_heading(text):
        return True

    return bool(
        bold
        and size >= median_font_size + 2.0
        and len(text.split()) <= 18
    )


def _load_pdf_sections(path: str) -> List[Dict[str, Any]]:
    """Convert PDF lines into section-aware ordered blocks."""
    lines = _extract_pdf_lines(path)
    if not lines:
        return []

    sizes = [
        line["font_size"]
        for line in lines
        if line["font_size"] > 0
    ]
    median_size = sorted(sizes)[len(sizes) // 2] if sizes else 10.0

    blocks: List[Dict[str, Any]] = []
    current_section = "Document"
    current_level = 0
    current_content: List[str] = []
    current_page = None

    def flush() -> None:
        nonlocal current_content
        if current_content:
            blocks.append(
                {
                    "section": current_section,
                    "section_level": current_level,
                    "content": "\n".join(current_content),
                    "page": current_page,
                }
            )
            current_content = []

    for line in lines:
        if _is_pdf_heading(line, median_size):
            flush()
            current_section = _normalize_heading(line["text"])
            marker = current_section.split(maxsplit=1)[0].rstrip(".")
            current_level = marker.count(".") + 1 if marker and marker[0].isdigit() else 1
            blocks.append(
                {
                    "section": current_section,
                    "section_level": current_level,
                    "content": current_section,
                    "page": line["page"],
                }
            )
            current_page = line["page"]
        else:
            current_content.append(line["text"])
            current_page = line["page"]

    flush()
    return blocks


# ---------------------------------------------------------------------------
# Plain extraction helpers
# ---------------------------------------------------------------------------


def _load_pdf(path: str) -> str:
    """Extract PDF text while preserving page boundaries."""
    pages: List[str] = []
    pdf = fitz.open(path)

    try:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text")
            if text and text.strip():
                pages.append(f"[Page {page_number}]\n{text.strip()}")
    finally:
        pdf.close()

    return "\n\n".join(pages)


def _load_docx(path: str) -> str:
    """Extract DOCX paragraphs and tables in original order."""
    sections = _load_docx_sections(path)
    return "\n".join(block["content"] for block in sections)


def _load_text(path: str) -> str:
    """Read a plain text file."""
    with open(path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


# ---------------------------------------------------------------------------
# Semantic chunking
# ---------------------------------------------------------------------------


def _split_large_section(section: str, content: str) -> List[str]:
    """Split a large section while retaining the section heading and overlap."""
    prefix = f"Section: {section}\n" if section != "Document" else ""

    available_size = max(100, CHUNK_SIZE - len(prefix))
    overlap = min(CHUNK_OVERLAP, max(0, available_size // 4))

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=available_size,
        chunk_overlap=overlap,
        separators=[
            "\n\n",
            "\n",
            ". ",
            "; ",
            ", ",
            " ",
            "",
        ],
    )

    return [
        prefix + piece
        for piece in splitter.split_text(content)
        if piece.strip()
    ]


def _chunk_sections(
    sections: List[Dict[str, Any]],
    source_name: str,
    source_type: str,
) -> List[Document]:
    """Turn ordered sections into coherent chunks with overlap for large sections."""
    if not sections:
        return []

    # Combine consecutive blocks belonging to the same section.
    grouped: List[Dict[str, Any]] = []
    for block in sections:
        section = block["section"]
        if grouped and grouped[-1]["section"] == section:
            grouped[-1]["contents"].append(block["content"])
        else:
            grouped.append(
                {
                    "section": section,
                    "section_level": block.get("section_level", 0),
                    "page": block.get("page"),
                    "contents": [block["content"]],
                }
            )

    documents: List[Document] = []
    chunk_index = 0

    for group in grouped:
        section = group["section"]
        section_text = "\n".join(
            item for item in group["contents"] if item.strip()
        ).strip()

        if not section_text:
            continue

        if len(section_text) <= CHUNK_SIZE:
            chunk_texts = [
                (
                    f"Section: {section}\n{section_text}"
                    if section != "Document"
                    else section_text
                )
            ]
        else:
            chunk_texts = _split_large_section(
                section=section,
                content=section_text,
            )

        for chunk_text in chunk_texts:
            metadata = {
                "source": source_name,
                "source_type": source_type,
                "section": section,
                "section_level": group["section_level"],
                "chunk_index": chunk_index,
                "chunking_strategy": "section_aware_overlap",
            }

            if group.get("page") is not None:
                metadata["page"] = group["page"]

            documents.append(
                Document(
                    page_content=chunk_text,
                    metadata=metadata,
                )
            )
            chunk_index += 1

    return documents


def _chunk_text(
    text: str,
    source_name: str,
    source_type: str,
) -> List[Document]:
    """Fallback recursive chunking for unstructured TXT content."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", "; ", " ", ""],
    )

    chunks = splitter.create_documents(
        [text],
        metadatas=[
            {
                "source": source_name,
                "source_type": source_type,
                "chunking_strategy": "recursive_overlap",
            }
        ],
    )

    documents: List[Document] = []
    for index, chunk in enumerate(chunks):
        metadata = dict(chunk.metadata)
        metadata["chunk_index"] = index
        documents.append(
            Document(
                page_content=chunk.page_content,
                metadata=metadata,
            )
        )

    return documents


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def load_and_split_document(path: str) -> List[Document]:
    """
    Extract and chunk PDF, DOCX, or TXT files.

    PDF/DOCX use section-aware chunking. Large sections are split with the
    configured CHUNK_OVERLAP while retaining the parent heading in every
    child chunk. TXT retains the existing recursive-overlap strategy.
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(f"File not found: {path}")

    extension = os.path.splitext(path)[1].lower()
    source_name = os.path.basename(path)

    if extension == ".pdf":
        source_type = "pdf"
        sections = _load_pdf_sections(path)
        documents = _chunk_sections(sections, source_name, source_type)

        if not documents:
            documents = _chunk_text(
                _load_pdf(path),
                source_name,
                source_type,
            )

    elif extension == ".docx":
        source_type = "docx"
        sections = _load_docx_sections(path)
        documents = _chunk_sections(sections, source_name, source_type)

    elif extension == ".txt":
        source_type = "transcript"
        documents = _chunk_text(
            _load_text(path),
            source_name,
            source_type,
        )

    else:
        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    if not documents:
        logger.warning(
            "No text/chunks extracted from %s",
            source_name,
        )
        return []

    logger.info(
        "Processed %s -> %d chunks | strategy=%s",
        source_name,
        len(documents),
        "section_aware_overlap" if extension in {".pdf", ".docx"} else "recursive_overlap",
    )

    for document in documents:
        logger.info(
            "CHUNK %s | section=%s | chars=%d",
            document.metadata.get("chunk_index"),
            document.metadata.get("section"),
            len(document.page_content),
        )

    return documents
