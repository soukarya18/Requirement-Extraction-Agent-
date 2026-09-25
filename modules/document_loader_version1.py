from __future__ import annotations

import os
from typing import List

import fitz
from docx import Document as DocxDocument

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from modules.config import CHUNK_OVERLAP, CHUNK_SIZE, get_logger


logger = get_logger("document_loader")


# ─── Extraction ───────────────────────────────────────────────────────────────


def _load_pdf(path: str) -> str:
    """Extract text from a PDF using PyMuPDF."""
    pages: List[str] = []

    pdf = fitz.open(path)

    try:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text")

            if text and text.strip():
                pages.append(
                    f"[Page {page_number}]\n{text.strip()}"
                )

    finally:
        pdf.close()

    return "\n\n".join(pages)


def _load_docx(path: str) -> str:
    """Extract paragraphs and tables from a DOCX document."""
    doc = DocxDocument(path)

    parts: List[str] = []

    # Paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in doc.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(cells)

            if row_text.strip():
                parts.append(row_text)

    return "\n".join(parts)


def _load_text(path: str) -> str:
    """Read a plain text transcript."""
    with open(
        path,
        "r",
        encoding="utf-8",
        errors="ignore",
    ) as file:
        return file.read()


# ─── Chunking ─────────────────────────────────────────────────────────────────


def _chunk_text(
    text: str,
    source_name: str,
    source_type: str,
) -> List[Document]:
    """Split text into LangChain Documents."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=[
            "\n\n",
            "\n",
            ". ",
            " ",
            "",
        ],
    )

    chunks = splitter.create_documents(
        [text],
        metadatas=[
            {
                "source": source_name,
                "source_type": source_type,
            }
        ],
    )

    documents: List[Document] = []

    for index, chunk in enumerate(chunks):
        metadata = dict(chunk.metadata)

        metadata["chunk_index"] = index

        documents.append(
            Document(
                page_content=chunk.page_content,
                metadata=metadata,
            )
        )

    return documents


# ─── Public API ───────────────────────────────────────────────────────────────


def load_and_split_document(
    path: str,
) -> List[Document]:
    """
    Extract and chunk a PDF, DOCX, or TXT file.

    Returns:
        List[Document]
    """
    if not os.path.isfile(path):
        raise FileNotFoundError(
            f"File not found: {path}"
        )

    extension = os.path.splitext(path)[1].lower()

    source_name = os.path.basename(path)

    if extension == ".pdf":
        text = _load_pdf(path)
        source_type = "pdf"

    elif extension == ".docx":
        text = _load_docx(path)
        source_type = "docx"

    elif extension == ".txt":
        text = _load_text(path)
        source_type = "transcript"

    else:
        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    if not text.strip():
        logger.warning(
            "No text extracted from %s",
            source_name,
        )
        return []

    documents = _chunk_text(
        text=text,
        source_name=source_name,
        source_type=source_type,
    )

    logger.info(
        "Processed %s -> %d chunks",
        source_name,
        len(documents),
    )

    return documents