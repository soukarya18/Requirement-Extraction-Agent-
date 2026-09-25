"""Render estimation results to Markdown and a downloadable DOCX report."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from estimation_generation.scoring import FactorScore


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_cell_border(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def _write_cell(cell, text: str, *, bold: bool = False, align=None) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _set_cell_border(cell)


def _markdown(factors: list[FactorScore], final_score: float, brs_name: str) -> str:
    lines = [
        "# BRD Delivery Estimation",
        "",
        f"**BRD:** {brs_name}",
        "",
        "## Scoring Summary",
        "",
        "Scores increase with delivery effort or risk. For requirement clarity, a higher score means greater ambiguity or incompleteness.",
        "",
        "| Factor | Weight | Score | Weighted contribution |",
        "| --- | ---: | ---: | ---: |",
    ]
    for factor in factors:
        lines.append(f"| {factor.label} | {factor.weight:.0%} | {factor.score} | {factor.score * factor.weight:.2f} |")
    lines.extend(["", "## Final Complexity Score", "", f"**{final_score:.2f} / 5.00**", ""])
    for factor in factors:
        lines.extend([f"## {factor.label}", "", f"**Assigned score:** {factor.score} / 5", "", f"**Justification:** {factor.justification}", "", "**Evidence observations:**"])
        lines.extend([f"- {observation}" for observation in factor.evidence_observations] or ["- No specific observation returned."])
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def write_estimation_report(output_dir: Path, brs_name: str, factors: list[FactorScore], final_score: float) -> tuple[Path, Path]:
    """Write matching Markdown and DOCX estimation artifacts."""
    output_dir.mkdir(parents=True, exist_ok=True)
    markdown_path = output_dir / "generated_estimation.md"
    docx_path = output_dir / "generated_estimation.docx"
    markdown_path.write_text(_markdown(factors, final_score, brs_name), encoding="utf-8")

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("BRD Delivery Estimation")
    title_run.font.name = "Aptos Display"
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph(f"BRD: {brs_name}")
    document.add_paragraph("This report presents evidence-grounded complexity scores and the calculated weighted complexity score. Scores increase with delivery effort or risk; for requirement clarity, a higher score means greater ambiguity or incompleteness.")

    document.add_heading("Scoring Summary", level=1)
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("Factor", "Weight", "Score", "Weighted contribution")):
        _set_cell_shading(cell, "1F4E78")
        _write_cell(cell, value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, factor in enumerate(factors):
        cells = table.add_row().cells
        if index % 2:
            for cell in cells:
                _set_cell_shading(cell, "EAF2F8")
        _write_cell(cells[0], factor.label)
        _write_cell(cells[1], f"{factor.weight:.0%}", align=WD_ALIGN_PARAGRAPH.CENTER)
        _write_cell(cells[2], str(factor.score), align=WD_ALIGN_PARAGRAPH.CENTER)
        _write_cell(cells[3], f"{factor.score * factor.weight:.2f}", align=WD_ALIGN_PARAGRAPH.CENTER)
    result = document.add_paragraph()
    result.add_run("Final complexity score: ").bold = True
    result.add_run(f"{final_score:.2f} / 5.00")

    document.add_heading("Factor Justifications", level=1)
    for factor in factors:
        document.add_heading(factor.label, level=2)
        score_paragraph = document.add_paragraph()
        score_paragraph.add_run("Assigned score: ").bold = True
        score_paragraph.add_run(f"{factor.score} / 5")
        document.add_paragraph(factor.justification)
        evidence_heading = document.add_paragraph()
        evidence_heading.add_run("Evidence observations").bold = True
        for observation in factor.evidence_observations or ["No specific observation returned."]:
            document.add_paragraph(observation, style="List Bullet")
    document.save(docx_path)
    return docx_path, markdown_path
